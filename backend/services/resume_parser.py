import io
import re
import zipfile
from pathlib import Path
from typing import Optional, Tuple

import pdfplumber
from docx import Document
import PyPDF2

from backend.utils.file_utils import(
    FileParsingError, 
    TextExtractionError, 
    log_error, 
    log_warning, 
    log_info, 
    with_fallback
)

from backend.core.config import (
    MAX_FILE_SIZE_BYTES,
    MAX_FILE_SIZE_MB, 
    SUPPORTED_MIME_TYPES
)

MAX_PDF_PAGES = 25

class FileValidationError(Exception):
    def __init__(self, message: str, user_message: Optional[str] = None):
        super().__init__(message)
        self.user_message = user_message or message

def sanitize_filename(filename: str) -> str:
    """Sanitize uploaded filename to prevent directory traversal."""
    if not filename:
        return "resume"
    clean = Path(filename).name
    # Keep only safe alphanumeric, dots, dashes, underscores
    clean = re.sub(r'[^a-zA-Z0-9_\-\.]', '_', clean)
    return clean or "resume"

def validate_file(file_data: bytes, filename: str) -> Tuple[bool, str, Optional[str]]:
    file_size_bytes = len(file_data)
    if file_size_bytes == 0:
        return False, 'Uploaded file is empty. Please check the file and try again.', None

    if file_size_bytes > MAX_FILE_SIZE_BYTES:
        size_mb = file_size_bytes / (1024 * 1024)
        return False, (
            f'File size ({size_mb:.2f} MB) exceeds the maximum allowed size of {MAX_FILE_SIZE_MB} MB. '
            'Please upload a smaller file or compress your resume.'
        ), None

    # Check extension
    clean_name = sanitize_filename(filename)
    ext = ''
    if '.' in clean_name:
        ext = '.' + clean_name.rsplit('.', 1)[-1].lower()

    if ext == '.doc':
        return False, (
            'Legacy .doc format is not supported. '
            'Please save or convert your document to .docx or .pdf and try again.'
        ), None

    # Magic byte verification
    if file_data.startswith(b'%PDF-'):
        mime_type = 'application/pdf'
    elif file_data.startswith(b'PK\x03\x04'):
        # Validate that this zip is indeed a valid Word docx
        try:
            with zipfile.ZipFile(io.BytesIO(file_data)) as zf:
                namelist = set(zf.namelist())
                if '[Content_Types].xml' in namelist or any(n.startswith('word/') for n in namelist):
                    mime_type = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
                else:
                    return False, (
                        'The uploaded file is a ZIP archive, but not a valid Word document (.docx). '
                        'Please upload a genuine .docx or .pdf file.'
                    ), None
        except zipfile.BadZipFile:
            return False, 'The uploaded .docx file is corrupted and cannot be read.', None
    else:
        mime_type = 'application/octet-stream'

    if mime_type not in SUPPORTED_MIME_TYPES:
        return False, (
            'Unsupported or invalid file format. '
            'Please upload a valid PDF (.pdf) or Word document (.docx).'
        ), None

    return True, '', SUPPORTED_MIME_TYPES[mime_type]


def _extract_pdf_hyperlinks(file_data: bytes) -> str:
    urls = []
    try:
        reader = PyPDF2.PdfReader(io.BytesIO(file_data))
        for page in reader.pages:
            if '/Annots' not in page:
                continue
            for annot_ref in page['/Annots']:
                try:
                    annot = annot_ref.get_object()
                    if annot.get('/Subtype') != '/Link':
                        continue
                    action = annot.get('/A', {})
                    uri = action.get('/URI', '')
                    if uri and isinstance(uri, (str, bytes)):
                        if isinstance(uri, bytes):
                            uri = uri.decode('utf-8', errors='ignore')
                        uri = uri.strip()
                        if uri.startswith('http'):
                            urls.append(uri)
                except Exception:
                    pass
    except Exception:
        pass
    return '\n'.join(urls)


def _check_pdf_protection_and_pages(file_data: bytes):
    """Check for password protection and excessive page counts."""
    try:
        reader = PyPDF2.PdfReader(io.BytesIO(file_data))
        if reader.is_encrypted:
            raise FileParsingError(
                'The PDF is password-protected or encrypted. '
                'Please remove password protection and re-upload your resume.'
            )
        page_count = len(reader.pages)
        if page_count > MAX_PDF_PAGES:
            raise FileParsingError(
                f'PDF exceeds the maximum allowed length of {MAX_PDF_PAGES} pages '
                f'({page_count} pages detected). Please upload a standard resume.'
            )
    except FileParsingError:
        raise
    except Exception:
        # PyPDF2 couldn't parse headers; fall through to pdfplumber
        pass


def _extract_pdf_with_pdfplumber(file_data: bytes) -> str:
    text = ''
    try:
        with pdfplumber.open(io.BytesIO(file_data)) as pdf:
            if len(pdf.pages) > MAX_PDF_PAGES:
                raise FileParsingError(
                    f'PDF exceeds the maximum allowed length of {MAX_PDF_PAGES} pages '
                    f'({len(pdf.pages)} pages detected).'
                )
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    text += page_text + '\n'
    except Exception as e:
        if "password" in str(e).lower():
            raise FileParsingError(
                'The PDF is password-protected or encrypted. '
                'Please remove password protection and re-upload your resume.'
            ) from e
        raise

    if not text.strip():
        raise TextExtractionError(
            'pdfplumber extracted no text',
            user_message='No text could be extracted from the PDF.'
        )
    
    hyperlinks = _extract_pdf_hyperlinks(file_data)
    if hyperlinks:
        text = text.strip() + '\n' + hyperlinks

    return text.strip()


def _extract_pdf_with_pypdf2(file_data: bytes) -> str:
    text = ''
    reader = PyPDF2.PdfReader(io.BytesIO(file_data))
    if reader.is_encrypted:
        raise FileParsingError(
            'The PDF is password-protected or encrypted. '
            'Please remove password protection and re-upload your resume.'
        )
    if len(reader.pages) > MAX_PDF_PAGES:
        raise FileParsingError(
            f'PDF exceeds the maximum allowed length of {MAX_PDF_PAGES} pages '
            f'({len(reader.pages)} pages detected).'
        )

    for page in reader.pages:
        page_text = page.extract_text()
        if page_text:
            text += page_text + '\n'

    if not text.strip():
        raise TextExtractionError(
            'PyPDF2 extracted no text',
            user_message='No text could be extracted from the PDF.'
        )

    hyperlinks = _extract_pdf_hyperlinks(file_data)
    if hyperlinks:
        text = text.strip() + '\n' + hyperlinks

    return text.strip()


def extract_text_from_pdf(file_data: bytes) -> str:
    _check_pdf_protection_and_pages(file_data)

    try: 
        result, used_fallback = with_fallback(
            _extract_pdf_with_pdfplumber, 
            _extract_pdf_with_pypdf2, 
            file_data, 
            log_fallback=True
        )
        if used_fallback:
            log_info('PDF extraction succeeded using the PyPDF2 fallback', context='resume_parser')

        # Check for scanned / image-only PDFs (no or minimal alphanumeric characters)
        alphanumeric_count = len(re.findall(r'[a-zA-Z0-9]', result))
        if alphanumeric_count < 20:
            raise FileParsingError(
                'No readable text could be found in the PDF. '
                'The file appears to be a scanned image or empty. '
                'Please upload a PDF containing selectable text or a DOCX document.'
            )

        return result
        
    except FileParsingError:
        raise
    except Exception as e:
        log_error(e, context='extract_text_from_pdf')
        raise FileParsingError(
            'Failed to extract text from PDF. '
            'The PDF may be corrupted, password-protected, or contain only scanned images. '
            'Please ensure it contains selectable text.'
        ) from e

    

def extract_text_from_docx(file_data: bytes) -> str:
    try:
        doc = Document(io.BytesIO(file_data))
        text_parts = []

        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                text_parts.append(paragraph.text)

        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        text_parts.append(cell.text)

        text = '\n'.join(text_parts)

        if not text.strip():
            raise FileParsingError(
                'No text could be extracted from the document. '
                'The document may be empty or corrupted.'
            )
        
        try:
            for rel in doc.part.rels.values():
                if 'hyperlink' in rel.reltype.lower():
                    url = rel._target
                    if isinstance(url, str) and url.startswith('http'):
                        text += '\n' + url
        except Exception:
            pass

        log_info(f'Extracted {len(text)} chars from DOCX', context='resume_parser')
        return text.strip()

    except FileParsingError:
        raise   # Re-raise unchanged — don't wrap in another FileParsingError

    except Exception as e:
        log_error(e, context='extract_text_from_docx')
        raise FileParsingError(
            'Failed to extract text from DOCX. '
            'The document may be corrupted or in an unsupported format. '
            'Please try re-saving or converting to PDF.'
        ) from e

def extract_text_from_doc(file_data: bytes) -> str:
    raise FileParsingError(
        'Legacy .doc format is not supported. '
        'Please convert your document to .docx or .pdf and try again. '
        'You can convert using Microsoft Word, Google Docs, or online tools.'
    )

def extract_text(file_data: bytes, file_type: str) -> str:
    if file_type == 'pdf':
        return extract_text_from_pdf(file_data)
    elif file_type == 'docx':
        return extract_text_from_docx(file_data)
    elif file_type == 'doc':
        return extract_text_from_doc(file_data)
    else:
        raise FileValidationError(
            f'Invalid file type: {file_type}. Supported types are: PDF (.pdf) and Word (.docx)'
        )

    
def parse_resume_file(file_data: bytes, filename: str) -> Tuple[str, dict]:
    clean_filename = sanitize_filename(filename)
    log_info(f'parsing file :{clean_filename}', context='parse_resume_file')

    # phase 01: validate file
    try:
        is_valid, error_msg, file_type = validate_file(file_data, clean_filename)
        if not is_valid:
            log_warning(f'validation failed for file {clean_filename}', context='parse_resume_file')
            raise FileValidationError(error_msg)
    
    except FileValidationError:
        raise 

    except Exception as e:
        log_error(e, context='parse_resume_file_validation')
        raise FileValidationError(
            'Could not validate the uploaded file. Please ensure it is a valid PDF or DOCX.'
        ) from e
    
    # phase 02: extraction of file
    try:
        text = extract_text(file_data, file_type)
        log_info(f'Extracted {len(text)} chars from {clean_filename}', context='parse_resume_file')

    except FileParsingError:
        raise   # Re-raise unchanged

    except Exception as e:
        log_error(e, context='parse_resume_file_extraction')
        raise FileParsingError(
            'An unexpected error occurred while processing the file. '
            'Please try again or contact support if the problem persists.'
        ) from e

    metadata = {
        'filename':        clean_filename,
        'file_type':       file_type,
        'file_size_bytes': len(file_data),
        'text_length':     len(text),
        'success':         True,
    }
    return text, metadata

